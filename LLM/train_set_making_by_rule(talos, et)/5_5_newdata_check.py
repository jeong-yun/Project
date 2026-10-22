"""[5.5 기존 학습데이터와 비교] 신규 학습데이터(08_dataset)를 기준 버전(기존 학습데이터)과 비교해 샘플마다 등급을 매긴다.  (화면 6)

    python 5_0_run_supervised.py compare --base "26 상반기" --dev ips --dev waf     ← 화면 6 의 '비교' 가 이렇게 실행

등급 (앞 등급이 우선. 표준모델 v26.03.03 2단계 08·10 과 같은 기준)
  제거  신규 내부 중복(정규화 payload = 소문자·공백 정리·숫자→0, 처음 한 줄만 남김) 또는 기준 데이터에 같은 payload
  A     기준 모델(기준 데이터로 학습한 RF·LGBM, 표준모델 피처 + 신규 후보 9개) 중 하나라도 틀림 = 기준 데이터에 부족한 데이터
  B     피처 패턴(43개 피처가 각각 켜졌는지의 조합)이 기준 데이터에 없음, 모델은 맞힘 = 새 패턴
  C     패턴도 기준 데이터에 있고 모델도 맞힘 = 이미 학습된 패턴
이미 비교한 데이터
  - 샘플별 피처는 payload 지문으로 캐시한다 → '새로 생긴 데이터만 비교' 는 새 샘플(새 룰·고친 샘플)의 피처만 계산한다.
  - 기준 버전까지 확정된 버전들이 이미 비교한 샘플은 빼고 비교한다 (계속 들어오는 새 룰만 보도록, 화면에서 다시 넣을 수 있음).
결과 (11_supervised/compared/<기준>/)
  <장비>_rows.parquet  샘플별 등급·근거      <장비>_table.json  화면 선별용 (표준 라이브러리로 읽음)
  <장비>_summary.json  화면 요약             <장비>_report.xlsx 엑셀 (요약 표 + 예시)
"""

import hashlib
import importlib.util
import json
import time
from pathlib import Path

import numpy as np
import pandas as pd

BASE = Path(__file__).resolve().parent


def _load(name, file):
    spec = importlib.util.spec_from_file_location(name, BASE / file)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


C = _load("sup_common", "5_0a_common.py")
F = _load("sup_features", "5_0b_features.py")
CHUNK = 4000                          # 피처 계산 묶음 (진행률·중지·캐시 단위)
FLAG_BITS = {"py_ua": 1, "base_conflict": 2, "collision": 4, "holdout": 8, "allzero": 16, "new_pat": 32,
             "wrong_rf": 64, "wrong_lg": 128}
Stopped = C.Stopped


class Ctx:
    """진행 보고 (실행기가 상태 파일·중지 요청을 붙인다. 따로 부를 때는 출력만)."""

    def __init__(self, pool=None, task=None):
        self.pool, self.task = pool, task          # 피처 계산 프로세스 묶음 (없으면 이 프로세스에서)

    def step(self, key, text):
        print(f"[{time.strftime('%H:%M:%S')}] {text}", flush=True)

    def progress(self, done, total=None, eta_s=None, unit="행"):
        pass

    def check(self):
        pass

    def log(self, msg):
        print(msg, flush=True)


def bin_key(F_: pd.DataFrame) -> np.ndarray:
    """피처 패턴 = 피처마다 켜졌는지(>0)의 조합 → uint64."""
    return pd.util.hash_pandas_object(F_.gt(0).astype(np.int8), index=False).to_numpy()


# ── 기준 데이터 ──────────────────────────────────────────────────────────

def load_base(v, ver, dev, ctx):
    """기준 학습 파일 → 피처(표준모델 + 신규 후보 9) · label · payload 지문 · 흔적. parquet 캐시 (파일이 바뀌면 다시)."""
    csv_p, _ = C.version_files(v, ver, dev)
    st = csv_p.stat()
    key = {"src": str(csv_p), "size": st.st_size, "mtime": st.st_mtime, "v2": F.V2_VERSION, "fmt": 1}
    out = C.CACHE / "base" / f"{ver['name']}_{dev}.parquet"
    meta = out.with_suffix(".json")
    if out.exists() and C.read_json(meta) == key:
        return pd.read_parquet(out)
    header = list(pd.read_csv(csv_p, nrows=0).columns)
    miss = [c for c in ("method", "payload", "label") if c not in header]
    if miss:
        raise RuntimeError(f"{csv_p.name} 에 열 {miss} 이 없습니다 (학습 파일은 method, payload, label, 피처 열)")
    feats = sorted(c for c in header if c.startswith(f"{dev}_payload_"))
    parts, t0, done = [], time.monotonic(), 0
    est = max(1, int(st.st_size / 640))                # 행 수 추정 (진행률용, 기존 학습 파일 평균 약 640바이트/행)
    for ch in pd.read_csv(csv_p, usecols=["method", "payload", "label"] + feats, dtype={c: "float32" for c in feats},
                          chunksize=40000, low_memory=False):
        ctx.check()
        p = ch["payload"].fillna("").astype(str)
        part = pd.concat([ch[["method", "label"] + feats].reset_index(drop=True),
                          F.generate_v2(p, dev).astype("float32").reset_index(drop=True)], axis=1)
        part["h_exact"] = F.hash_series(p)
        part["h_norm"] = F.hash_series(p.map(F.norm_payload))
        part["ua"] = F.ua_category(p).to_numpy()
        part["hdr"] = F.header_count(p).astype("int16").to_numpy()
        part["plen"] = p.str.len().astype("int32").to_numpy()
        parts.append(part)
        done += len(ch)
        rate = done / max(1e-6, time.monotonic() - t0)
        ctx.progress(done, max(est, done), (max(est, done) - done) / rate)
    B = pd.concat(parts, ignore_index=True)
    for c in ("method", "label", "ua"):
        B[c] = B[c].astype("category")
    out.parent.mkdir(parents=True, exist_ok=True)
    B.to_parquet(out)
    C.write_json(meta, key)
    return B


def base_counts(B):
    return {"rows": int(len(B)), "label": {k: int(n) for k, n in B["label"].value_counts().items()},
            "methods": {str(k): int(n) for k, n in B["method"].value_counts().head(30).items()}}


def base_models(ver, dev, B, cols, mode, ctx):
    """기준 데이터로 학습한 모델 {(구성, rf|lgbm): 모델}. 기준 데이터가 같으면 저장해 둔 모델을 쓴다."""
    import joblib
    d = C.CACHE / "models" / ver["name"]
    d.mkdir(parents=True, exist_ok=True)
    key = C.read_json(C.CACHE / "base" / f"{ver['name']}_{dev}.json", {})
    y = (B["label"] == "anomalies").to_numpy()
    out, todo = {}, []
    for cfg, cs in cols.items():
        for m in ("rf", "lgbm"):
            p = d / f"{dev}_{cfg}_{m}.joblib"
            meta = C.read_json(p.with_suffix(".json"), {})
            if mode != "all" and p.exists() and meta == dict(key, cols=cs):
                out[(cfg, m)] = joblib.load(p)
            else:
                todo.append((cfg, m, cs, p))
    t0 = time.monotonic()
    for k, (cfg, m, cs, p) in enumerate(todo):
        ctx.check()
        ctx.progress(k, len(todo), None, unit="모델")
        model = C.make_model(m).fit(B[cs].to_numpy(dtype=np.float32), y)
        joblib.dump(model, p, compress=3)
        C.write_json(p.with_suffix(".json"), dict(key, cols=cs))
        out[(cfg, m)] = model
        ctx.log(f"  기준 모델 {cfg} {m}: 학습 {time.monotonic() - t0:.0f}초, 파일 {p.stat().st_size / 2**20:.0f}MB")
    return out


# ── 신규 학습데이터 + 피처 ───────────────────────────────────────────────

def read_samples(dev, dates=None):
    """08_dataset 의 모든(또는 고른) 날짜 → 샘플 표 (날짜·줄 순서), 날짜별 로그 상태."""
    cols, stats = {c: [] for c in C.SAMPLE_COLS}, {}
    for d in dates or C.dataset_dates():
        got, stat = C.read_dataset(d, dev)
        if stat is None:
            continue
        stats[d] = stat
        for k in cols:
            cols[k].extend(got[k])
    S = pd.DataFrame(cols)
    S["ck"] = [C.sample_key(a, b, c) for a, b, c in zip(S["label"], S["rule_id"], S["payload_raw"])]
    S["pk"] = [C.payload_key(p) for p in S["payload_raw"]]
    return S, stats


def _feature_dir(dev, code):
    return C.CACHE / "features" / dev / F.code_hash(code)


def load_feature_cache(d):
    parts = sorted(d.glob("part-*.parquet"))
    if not parts:
        return None
    df = pd.concat([pd.read_parquet(p) for p in parts], ignore_index=True)
    return df.drop_duplicates("pk", keep="last").set_index("pk")


def features_for(S, dev, code, mode, ctx):
    """샘플 표의 각 줄 피처 (캐시에 없는 payload 만 계산). mode=all 이면 모두 다시 계산."""
    d = _feature_dir(dev, code)
    for old in d.parent.glob("*"):                       # 피처 파일·신규 후보가 바뀌어 쓰지 않는 캐시
        if old != d and old.is_dir():
            for f in old.glob("*"):
                f.unlink()
            old.rmdir()
    d.mkdir(parents=True, exist_ok=True)
    if mode == "all":
        for f in d.glob("part-*.parquet"):
            f.unlink()
    cache = load_feature_cache(d)
    have = set(cache.index) if cache is not None else set()
    todo = S.loc[~S["pk"].isin(have), ["pk", "payload_raw"]].drop_duplicates("pk")
    n, t0, done = len(todo), time.monotonic(), 0
    ctx.log(f"  피처: 전체 payload {S['pk'].nunique():,}개 중 새로 계산 {n:,}개 (캐시 {len(have):,}개)")
    stamp = time.strftime("%Y%m%d%H%M%S")
    chunks = [todo.iloc[s:s + CHUNK] for s in range(0, n, CHUNK)]

    def save(k, ch, f):
        f.insert(0, "pk", ch["pk"].to_numpy())
        f.to_parquet(d / f"part-{stamp}-{k:05d}.parquet", index=False)

    def tick(rows):
        nonlocal done
        done += rows
        rate = done / max(1e-6, time.monotonic() - t0)
        ctx.progress(done, n, (n - done) / rate if done else None)

    ctx.progress(0, n, None)
    if ctx.pool is not None and len(chunks) > 1:
        from concurrent.futures import FIRST_COMPLETED, wait
        pending, k = {}, 0
        width = max(2, getattr(ctx.pool, "_max_workers", 2) * 2)
        while k < len(chunks) or pending:
            while k < len(chunks) and len(pending) < width:
                ch = chunks[k]
                pending[ctx.pool.submit(ctx.task, dev, str(code), ch["payload_raw"].tolist())] = (k, ch)
                k += 1
            fin, _ = wait(list(pending), return_when=FIRST_COMPLETED)
            for fu in fin:
                kk, ch = pending.pop(fu)
                save(kk, ch, fu.result())
                tick(len(ch))
            try:
                ctx.check()
            except Stopped:
                for fu in pending:
                    fu.cancel()
                raise
    else:
        for k, ch in enumerate(chunks):
            ctx.check()
            save(k, ch, F.compute_chunk(dev, str(code), ch["payload_raw"].tolist()))
            tick(len(ch))
    if n:
        timing = C.read_json(C.CACHE / "timing.json", {})
        timing[f"feat_rows_per_sec_{dev}"] = round(n / max(1e-6, time.monotonic() - t0), 1)
        C.write_json(C.CACHE / "timing.json", timing)
    parts = sorted(d.glob("part-*.parquet"))
    cache = load_feature_cache(d)
    if len(parts) > 30:                                  # 작은 파일이 많아지면 하나로 합침
        cache.reset_index().to_parquet(d / f"part-{stamp}-all.parquet", index=False)
        for p in parts:
            p.unlink()
    return cache.loc[S["pk"].to_numpy()].reset_index(drop=True)


# ── 등급 ────────────────────────────────────────────────────────────────

def grade(S, Fz, B, models, cols, dev):
    """샘플 표 + 피처 → 등급·근거 열을 붙인 표 (S 와 같은 순서)."""
    y = S["label"].eq("anomalies").to_numpy()
    yB = B["label"].eq("anomalies").to_numpy()
    R = S.drop(columns=["payload_raw"]).copy()
    for (cfg, m), model in models.items():
        R[f"p_{m}{cfg}"] = model.predict_proba(Fz[cols[cfg]].to_numpy(dtype=np.float32))[:, 1].astype("float32")
    wrong_rf = (R["p_rf43"].to_numpy() >= 0.5) != y
    wrong_lg = (R["p_lgbm43"].to_numpy() >= 0.5) != y
    # 패턴 (43개 피처 on/off): 기준 데이터에 있는지, 기준에서는 반대 라벨만 있었는지
    patB, patN = bin_key(B[cols["43"]]), bin_key(Fz[cols["43"]])
    pc = pd.DataFrame({"p": patB, "y": yB}).groupby("p")["y"].agg(["sum", "count"])
    n_att = pd.Series(patN).map(pc["sum"]).fillna(0).to_numpy()
    n_all = pd.Series(patN).map(pc["count"]).fillna(0).to_numpy()
    n_norm = n_all - n_att
    new_pat = n_all == 0
    collision = ~new_pat & (np.where(y, n_att, n_norm) == 0)
    new34 = ~pd.Series(bin_key(Fz[cols["34"]])).isin(set(bin_key(B[cols["34"]]))).to_numpy()
    # 중복: 신규 내부(정규화, 처음 한 줄만 남김) / 기준 데이터에 같은 payload(정규화)
    hn = pd.Series(Fz["h_norm"].to_numpy())
    dup_int = hn.duplicated().to_numpy()
    bl = pd.DataFrame({"h": B["h_norm"].to_numpy(), "y": yB}).groupby("h")["y"].agg(["min", "max"])
    in_base = hn.isin(bl.index).to_numpy()
    bmin = hn.map(bl["min"]).to_numpy()
    bmax = hn.map(bl["max"]).to_numpy()
    base_conflict = in_base & ((bmin != y) | (bmax != y))
    g = np.select([dup_int | in_base, wrong_rf | wrong_lg, new_pat], [0, 1, 2], 3)
    R["grade"] = np.array(C.GRADES)[g]
    allzero = (Fz[cols["34"]].abs().sum(axis=1) == 0).to_numpy()
    R["pat"] = patN
    R["pat_id"] = pd.factorize(patN)[0].astype("int32")
    R["holdout"] = [C.holdout(r) for r in S["rule_id"]]
    for k, val in (("new_pat", new_pat), ("new_pat34", new34), ("wrong_rf", wrong_rf), ("wrong_lg", wrong_lg),
                   ("dup_int", dup_int), ("dup_exact_int", pd.Series(Fz["h_exact"].to_numpy()).duplicated().to_numpy()),
                   ("dup_base", in_base), ("dup_exact_base", pd.Series(Fz["h_exact"].to_numpy()).isin(set(B["h_exact"])).to_numpy()),
                   ("base_conflict", base_conflict), ("collision", collision), ("allzero", allzero & y)):
        R[k] = val
    for c in ("sigs", "ua", "py_ua", "hdr", "plen"):
        R[c] = Fz[c].to_numpy()
    # 신규 내부 라벨 상충 (정규화 payload 가 같은데 라벨이 다름)
    lab_by_h = pd.DataFrame({"h": hn, "y": y}).groupby("h")["y"].agg(["min", "max"])
    R["int_conflict"] = hn.map(lab_by_h["min"] != lab_by_h["max"]).to_numpy()
    return R


# ── 요약 (화면·엑셀 공용) ────────────────────────────────────────────────

def _cap(s, cap):
    return int(s.value_counts().clip(upper=cap).sum()) if len(s) else 0


def summarize(R, B, S, Fz, cols, dev, base, stats, scope, models_info):
    y = R["label"].eq("anomalies").to_numpy()
    yB = B["label"].eq("anomalies").to_numpy()
    lab = np.where(y, "공격", "정상")
    out = {"dev": dev, "base": base["name"], "rows": int(len(R)), "label": {"공격": int(y.sum()), "정상": int((~y).sum())},
           "created_at": C.now_iso(), "log_stats": stats, "scope": scope, "tables": {}, "kv": []}

    def kv(k, val, note=""):
        out["kv"].append([k, val, note])

    def table(name, df, note=""):
        out["tables"][name] = {"cols": [str(c) for c in df.columns], "rows": df.astype(object).where(df.notna(), None).values.tolist(),
                               "note": note}

    # 개요
    kv("신규 행 수 (비교 범위)", f"{len(R):,}", f"공격 {int(y.sum()):,} / 정상 {int((~y).sum()):,}")
    kv("날짜", ", ".join(f"{d} ({n:,})" for d, n in R["date"].value_counts().sort_index().items()))
    kv("출처 (rule_id 앞부분)", ", ".join(f"{k} {n:,}" for k, n in R["source"].value_counts().items()))
    kv("rule_id 수", f"{R['rule_id'].nunique():,}")
    kv("증강 샘플 비율", f"{(R['origin'] == 'augmented').mean():.1%}" if len(R) else "-")
    if scope.get("excluded"):
        kv("이미 비교한 데이터 (확정된 버전, 뺌)", f"{scope['excluded']:,}")
    nv = sorted(set(R.loc[R["cons"] == "", "date"]))
    kv("label 정합성 (4. 화면 '전체 검증')", "모든 날짜 검증됨" if not nv else f"검증 안 됨: {', '.join(nv)}",
       "'정합 (룰 기준)' 빼기·'불일치' 빼기는 검증된 날짜에만 적용됩니다" if nv else "")
    kv("기준 버전", f"{base['name']} · {len(B):,}행 (공격 {int(yB.sum()):,} / 정상 {int((~yB).sum()):,})")

    # 등급 × label (+ 패턴당 상한 적용 시 수)
    rows = []
    for gname in C.GRADES:
        m = (R["grade"] == gname).to_numpy()
        r = {"등급": C.GRADE_TEXT[gname], "공격": int((m & y).sum()), "정상": int((m & ~y).sum()), "합계": int(m.sum()),
             "패턴 수": int(R.loc[m, "pat"].nunique())}
        for cap in (20, 50):
            r[f"패턴당 {cap}건이면"] = sum(_cap(R.loc[m & (y == lv), "pat"], cap) for lv in (True, False))
        rows.append(r)
    t = pd.DataFrame(rows)
    t.loc[len(t)] = ["합계", int(y.sum()), int((~y).sum()), len(R), int(R["pat"].nunique()), "", ""]
    table("등급별", t, "제거 > A > B > C 순서로 먼저 맞는 등급. 패턴 = 43개 피처(표준모델 + 신규 후보 9)가 각각 켜졌는지의 조합")
    a = (R["grade"] == "A").to_numpy()
    wr, wl, newp = R["wrong_rf"].to_numpy(), R["wrong_lg"].to_numpy(), R["new_pat"].to_numpy()
    table("A 근거", pd.DataFrame([
        {"구분": "RF·LGBM 둘 다 틀림", "공격(미탐)": int((a & wr & wl & y).sum()), "정상(오탐)": int((a & wr & wl & ~y).sum())},
        {"구분": "RF만 틀림", "공격(미탐)": int((a & wr & ~wl & y).sum()), "정상(오탐)": int((a & wr & ~wl & ~y).sum())},
        {"구분": "LGBM만 틀림", "공격(미탐)": int((a & ~wr & wl & y).sum()), "정상(오탐)": int((a & ~wr & wl & ~y).sum())},
        {"구분": "└ 그중 새 패턴", "공격(미탐)": int((a & newp & y).sum()), "정상(오탐)": int((a & newp & ~y).sum())},
        {"구분": "└ 그중 기준에 있던 패턴", "공격(미탐)": int((a & ~newp & y).sum()), "정상(오탐)": int((a & ~newp & ~y).sum())}]),
        "기준 모델 = 기준 데이터로 학습한 RF·LGBM (표준모델 피처 + 신규 후보 9개, 0.5 기준)")
    for name, col in (("공격 유형별", "attack_type"), ("출처별", "source"), ("날짜별", "date")):
        ct = pd.crosstab(R[col], R["grade"]).reindex(columns=C.GRADES, fill_value=0)
        ct["합계"] = ct.sum(axis=1)
        table(name, ct.sort_values("합계", ascending=False).reset_index().rename(columns={col: name[:-1]}))

    # 품질
    q = [("신규 내부 완전 중복 (행)", int(R["dup_exact_int"].sum())),
         ("신규 내부 정규화 중복 (행, 숫자만 다른 것 포함)", int(R["dup_int"].sum())),
         ("신규 내부 라벨 상충 (정규화 payload 같음, 행)", int(R["int_conflict"].sum())),
         ("기준 데이터와 완전 중복 (행)", int(R["dup_exact_base"].sum())),
         ("기준 데이터와 정규화 중복 (행)", int(R["dup_base"].sum())),
         ("기준 데이터와 라벨 상충 (정규화 기준, 행)", int(R["base_conflict"].sum())),
         ("정상 라벨 + 강한 공격 시그니처", int(((~y) & R["sigs"].ne("").to_numpy()).sum())),
         ("공격 라벨 + 표준모델 피처 전부 0 + 시그니처 없음", int((R["allzero"].to_numpy() & R["sigs"].eq("").to_numpy()).sum())),
         ("'정합 (룰 기준)' 공격 (보류 옵션 대상, 검증된 날짜만)", int((y & R["cons"].eq("RULE").to_numpy()).sum())),
         ("python-requests UA 정상 (보류 옵션 대상)", int(((~y) & R["py_ua"].to_numpy()).sum())),
         ("'불일치' (4. 화면 검증, 검증된 날짜만)", int(R["cons"].eq("BAD").sum())),
         ("충돌 위험: 기준에서는 반대 라벨만 있던 패턴", int(R["collision"].sum()))]
    table("품질", pd.DataFrame(q, columns=["점검", "건수"]))
    cons = pd.crosstab(R["cons"].map(C.CONS_TEXT), lab)
    table("label 정합성", cons.reset_index().rename(columns={"cons": "label 정합성"}))

    # 기준 모델의 신규 탐지율·오탐율
    det = []
    for cfg, name in (("34", "표준모델 피처"), ("43", "+ 신규 후보 9 (등급 기준)")):
        for m in ("rf", "lgbm"):
            hit = R[f"p_{m}{cfg}"].to_numpy() >= 0.5
            det.append({"구성": name, "모델": m.upper(), "신규 공격 탐지율": round(float(hit[y].mean()), 4) if y.any() else None,
                        "신규 정상 오탐율": round(float(hit[~y].mean()), 4) if (~y).any() else None,
                        "미탐 수": int((y & ~hit).sum()), "오탐 수": int((~y & hit).sum())})
    table("기준 모델의 신규 탐지", pd.DataFrame(det), "기준 데이터 전체로 학습한 모델이 신규 데이터를 맞히는지 (오탐율이 높으면 신규 정상이 필요한 데이터)")

    # 데이터 흔적 (UA·헤더·길이) - 신규 vs 기준
    ua = pd.concat([pd.crosstab(R["ua"], lab, normalize="columns").add_prefix("신규 "),
                    pd.crosstab(B["ua"].astype(str), np.where(yB, "공격", "정상"), normalize="columns").add_prefix("기준 ")],
                   axis=1).fillna(0).round(4)
    table("흔적: User-Agent", ua.reset_index().rename(columns={"index": "UA", "ua": "UA", "row_0": "UA"}),
          "공격·정상의 UA 분포가 다르면 모델이 UA 로 구분하게 될 수 있음")
    hl = []
    for name, hdr, plen, mask in (("신규 공격", R["hdr"], R["plen"], y), ("신규 정상", R["hdr"], R["plen"], ~y),
                                  ("기준 공격", B["hdr"], B["plen"], yB), ("기준 정상", B["hdr"], B["plen"], ~yB)):
        hl.append({"구분": name, "헤더 줄 수 중앙값": float(np.median(hdr[mask])) if mask.any() else None,
                   "payload 길이 중앙값": float(np.median(plen[mask])) if mask.any() else None})
    table("흔적: 헤더·길이", pd.DataFrame(hl))

    # 표준모델 피처로 못 잡는 공격, 커버리지
    zero = R["allzero"].to_numpy()
    kv("표준모델 피처가 하나도 안 켜지는 공격", f"{int(zero.sum()):,}", "이 공격들은 표준모델 피처로는 정상과 구분할 근거가 없음")
    on = Fz[cols["34"]].gt(0).to_numpy()
    single = y & (on.sum(axis=1) == 1)
    if single.any():
        names = np.array(cols["34"])[on[single].argmax(axis=1)]
        sv = pd.Series(names).value_counts().head(10)
        table("표준모델 피처 1개만 켜지는 공격", pd.DataFrame({"켜진 피처": sv.index, "공격 수": sv.values}),
              "기존 미탐이 많이 몰리는 지점 (그 피처 하나로만 공격 판단)")
    kv("새 패턴 비율 (43개 피처 기준)", f"{R['new_pat'].mean():.1%}" if len(R) else "-",
       f"표준모델 피처 34개 기준 {R['new_pat34'].mean():.1%}" if len(R) else "")
    kv("기준에는 정상만 있던 패턴에 떨어지는 신규 공격", f"{int((y & R['collision'].to_numpy()).sum()):,}", "오탐·미탐 충돌 위험")

    # 평가 분할 (rule_id 지문 70/30)
    sp = pd.crosstab(np.where(R["holdout"], "holdout (30%)", "train (70%)"), lab)
    sp["rule_id 수"] = [R.loc[R["holdout"] == (k.startswith("holdout")), "rule_id"].nunique() for k in sp.index]
    table("평가 분할", sp.reset_index().rename(columns={"row_0": "분할", "index": "분할"}),
          "결합 학습을 평가할 때 holdout 룰은 학습에 넣지 않고 평가에만 쓰도록 표시 (지문으로 정해 새 룰이 와도 그대로)")
    out["models"] = models_info
    return out


def examples(R, S):
    """엑셀 예시 시트."""
    X = R.copy()
    X["payload"] = S["payload_raw"].str.slice(0, 1500).to_numpy()
    cols = ["date", "line", "rule_id", "method", "label", "attack_type", "origin", "cons", "grade", "sigs", "payload"]
    y = X["label"].eq("anomalies")
    sheets = {
        "예시_A 공격(룰별1건)": X[(X["grade"] == "A") & y].groupby("rule_id").head(1).head(300),
        "예시_A 정상(룰별1건)": X[(X["grade"] == "A") & ~y].groupby("rule_id").head(1).head(300),
        "예시_정상+공격시그니처": X[~y & X["sigs"].ne("")].head(200),
        "예시_정합(룰기준)": X[y & X["cons"].eq("RULE")].groupby("rule_id").head(1).head(200),
        "예시_표준피처 전부0 공격": X[X["allzero"]].groupby("rule_id").head(1).head(200),
        "예시_기준과 라벨상충": X[X["base_conflict"]].head(200),
    }
    return {k: v[cols] for k, v in sheets.items()}


def write_report(path, summary, ex):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    with pd.ExcelWriter(tmp, engine="openpyxl") as xw:
        pd.DataFrame(summary["kv"], columns=["항목", "값", "설명"]).to_excel(xw, sheet_name="요약", index=False)
        for name, t in summary["tables"].items():
            pd.DataFrame(t["rows"], columns=t["cols"]).to_excel(xw, sheet_name=name.replace(":", "")[:31], index=False)
        for name, d in ex.items():
            d = d.copy()
            for c in d.columns:
                if d[c].dtype == object:
                    d[c] = d[c].astype(str).str.replace(ILLEGAL_CHARACTERS_RE, "", regex=True)
            d.to_excel(xw, sheet_name=name[:31], index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E78")
            for col in ws.columns:
                w = max(len(str(c.value or "")) for c in list(col)[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, w * 1.2), 70)
    tmp.replace(path)


def write_table(path, R, base_name, dev, snapshot):
    """화면 선별용 표 (표준 라이브러리 json 으로 읽음). 열마다 값 목록."""
    dates = sorted(R["date"].unique().tolist())
    rules = pd.unique(R["rule_id"]).tolist()
    types = sorted(R["attack_type"].unique().tolist())
    srcs = sorted(R["source"].unique().tolist())
    idx = lambda col, vals: pd.Series(R[col]).map({v: i for i, v in enumerate(vals)}).astype(int).tolist()
    fl = np.zeros(len(R), dtype=np.int64)
    for k, bit in FLAG_BITS.items():
        fl |= np.where(R[k].to_numpy(dtype=bool), bit, 0)
    cons_num = {v: k for k, v in C.CONS_CODE.items()}
    data = {"base": base_name, "dev": dev, "snapshot": snapshot, "created_at": C.now_iso(), "n": int(len(R)),
            "dates": dates, "rules": rules, "types": types, "sources": srcs,
            "date": idx("date", dates), "line": R["line"].astype(int).tolist(),
            "g": pd.Series(R["grade"]).map({g: i for i, g in enumerate(C.GRADES)}).astype(int).tolist(),
            "y": R["label"].eq("anomalies").astype(int).tolist(), "pat": R["pat_id"].astype(int).tolist(),
            "rule": idx("rule_id", rules), "type": idx("attack_type", types), "src": idx("source", srcs),
            "llm": R["origin"].eq("llm").astype(int).tolist(),
            "cons": "".join(cons_num.get(c, "0") for c in R["cons"]), "fl": fl.tolist()}
    C.write_json(path, data, indent=None)


def snapshot_id(R):
    h = hashlib.sha1()
    for ck, g in zip(R["ck"], R["grade"]):
        h.update(f"{ck}{g}".encode())
    return h.hexdigest()[:12]


# ── 실행 ────────────────────────────────────────────────────────────────

def compare_dev(v, base, dev, mode, include_compared, ctx):
    code = C.feature_code(v, dev)
    ns = F.load_legacy(code)
    names34 = F.legacy_names(ns)
    ctx.step(f"{dev}:base", f"{dev.upper()} 기준 데이터 읽기 ({base['name']})")
    B = load_base(v, base, dev, ctx)
    cols34 = sorted(c for c in B.columns if c.startswith(f"{dev}_payload_") and c not in F.v2_cols(dev))
    if cols34 != names34:
        raise RuntimeError(f"{base['name']} {dev.upper()} 학습 파일의 피처 열({len(cols34)}개)이 표준모델 피처 파일의 피처"
                           f"({len(names34)}개)와 다릅니다 - 같은 표준모델 버전의 파일인지 확인하세요")
    cols = {"34": cols34, "43": cols34 + F.v2_cols(dev)}
    v_now = C.load_versions()                             # 기준 버전 행 수 (버전 표)
    for x in v_now["versions"]:
        if x["name"] == base["name"]:
            x.setdefault("counts", {})[dev] = base_counts(B)
    C.save_versions(v_now)
    ctx.step(f"{dev}:model", f"{dev.upper()} 기준 모델 (RF·LGBM, 처음 한 번 학습)")
    t0 = time.monotonic()
    models = base_models(base, dev, B, cols, mode, ctx)
    models_info = {"train_s": round(time.monotonic() - t0, 1)}
    ctx.step(f"{dev}:read", f"{dev.upper()} 신규 학습데이터 읽기")
    S, stats = read_samples(dev)
    scope = {"all": int(len(S)), "excluded": 0, "include_compared": include_compared}
    if not include_compared:
        done = C.considered_keys(v, base["name"], dev)
        if done:
            keep = ~S["ck"].isin(done)
            scope["excluded"] = int((~keep).sum())
            S = S.loc[keep].reset_index(drop=True)
    if S.empty:
        raise RuntimeError(f"{dev.upper()} 비교할 신규 학습데이터가 없습니다")
    ctx.step(f"{dev}:feat", f"{dev.upper()} 피처 계산 (새 샘플만)" if mode != "all" else f"{dev.upper()} 피처 계산 (전체)")
    Fz = features_for(S, dev, code, mode, ctx)
    ctx.step(f"{dev}:grade", f"{dev.upper()} 등급 매기기")
    R = grade(S, Fz, B, models, cols, dev)
    snap = snapshot_id(R)
    ctx.step(f"{dev}:save", f"{dev.upper()} 결과 저장")
    files = C.compared_files(base["name"], dev)
    files["rows"].parent.mkdir(parents=True, exist_ok=True)
    R.to_parquet(files["rows"], index=False)
    write_table(files["table"], R, base["name"], dev, snap)
    summ = summarize(R, B, S, Fz, cols, dev, base, stats, scope, models_info)
    summ["snapshot"] = snap
    summ["mode"] = mode
    write_report(files["report"], summ, examples(R, S))
    C.write_json(files["summary"], summ)
    g = R["grade"].value_counts()
    ctx.log(f"  {dev.upper()} {len(R):,}행: " + " · ".join(f"{k} {int(g.get(k, 0)):,}" for k in C.GRADES) + f" (스냅숏 {snap})")
    return summ


def run(base_name, devs, mode="new", include_compared=False, ctx=None):
    ctx = ctx or Ctx()
    v = C.load_versions()
    problems = C.setup_problems(v, base_name, devs)
    if problems:
        raise RuntimeError("; ".join(problems))
    base = C.get_version(v, base_name)
    return {dev: compare_dev(v, base, dev, mode, include_compared, ctx) for dev in devs}
